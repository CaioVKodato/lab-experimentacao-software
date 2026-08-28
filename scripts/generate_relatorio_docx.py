"""
Gera o relatório final Lab01 a partir do template .docx da disciplina.

Uso: python scripts/generate_relatorio_docx.py
"""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

BLACK = RGBColor(0, 0, 0)
BODY_FONT = "Arial"
BODY_SIZE = Pt(11)

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
TEMPLATE = ROOT / "docs" / "template" / "Template_Relatorio_Laboratorio.docx"
OUT_DOCX = DOCS / "relatorios" / "relatorio_lab01_final.docx"

SKIP_LINK_PATTERNS = [
    r"youtube",
    r"youtu\.be",
    r"www\.you",
    r"favor inserir a referência de vídeo",
]

SKIP_PATTERNS = [
    r"^ORIENTAÇÃO:",
    r"^Este documento é um MODELO",
    r"^Perguntas que esta seção deve responder",
    r"^Qual problema está sendo investigado",
    r"^Quais são as Questões de Pesquisa do enunciado",
    r"^Quais as hipóteses informais do grupo",
    r"^Quais RQs, métricas ou variáveis o grupo está propondo",
    r"^Sugestão: insira aqui o print",
]


def is_trap(text: str) -> bool:
    low = text.lower()
    return any(re.search(p, low) for p in SKIP_LINK_PATTERNS)


def should_skip(text: str) -> bool:
    t = text.strip()
    if not t:
        return False
    if is_trap(t):
        return True
    return any(re.search(p, t) for p in SKIP_PATTERNS)


def delete_paragraph(paragraph: Paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def fix_run_formatting(run, *, heading: bool = False) -> None:
    run.font.color.rgb = BLACK
    run.font.italic = False
    run.font.name = BODY_FONT
    if not heading:
        run.font.size = BODY_SIZE
        run.font.bold = False


def fix_paragraph_formatting(paragraph: Paragraph) -> None:
    heading = paragraph.style.name.startswith("Heading") or paragraph.style.name == "Title"
    for run in paragraph.runs:
        fix_run_formatting(run, heading=heading)


def fix_document_formatting(doc: Document) -> None:
    for para in doc.paragraphs:
        fix_paragraph_formatting(para)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    fix_paragraph_formatting(para)


def replace_text(paragraph: Paragraph, text: str) -> None:
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)
    fix_paragraph_formatting(paragraph)


def insert_paragraph_after(ref: Paragraph, text: str = "") -> Paragraph:
    new_p = OxmlElement("w:p")
    ref._element.addnext(new_p)
    para = Paragraph(new_p, ref._parent)
    if text:
        run = para.add_run(text)
        fix_run_formatting(run)
    return para


def insert_picture_after(ref: Paragraph, path: Path, width_cm: float) -> Paragraph:
    para = insert_paragraph_after(ref)
    run = para.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    return para


CONTENT_BLOCKS = [
    "intro",
    "contexto",
    "desafios",
    "decisoes",
    "ferramentas",
    "inovacoes",
    "coleta",
    "discussao",
    "conclusao",
]

CONTENT: dict[str, str] = {
    "intro": (
        "Este laboratório investiga características de repositórios open-source populares no GitHub. "
        "Coletamos os 1.000 repositórios com mais estrelas e respondemos sete questões de pesquisa "
        "sobre maturidade, contribuição externa, releases, atualização de código, linguagem primária, "
        "taxa de issues fechadas e o cruzamento dessas métricas por linguagem.\n\n"
        "Hipóteses informais (antes da análise):\n"
        "RQ01 — repositórios populares tendem a ser antigos; "
        "RQ02 — recebem muita contribuição externa via PRs; "
        "RQ03 — lançam releases com frequência, mas não em todos; "
        "RQ04 — são atualizados com frequência; "
        "RQ05 — concentram-se nas linguagens mais usadas no GitHub; "
        "RQ06 — quem usa o tracker de issues tem alta taxa de fechamento; "
        "RQ07 — linguagens do top-10 Octoverse têm mais PRs, releases e pushes recentes.\n\n"
        "Inovação do grupo (30%): RQ extra — popularidade (stars) se correlaciona com o percentual "
        "de issues fechadas? Hipótese: não; o detalhamento está na seção 3.6."
    ),
    "contexto": (
        "O Lab01 é o primeiro laboratório do semestre e inaugura o Kanban do grupo (GitHub Projects v2), "
        "mantido até o Lab05. Na S01 montamos o cliente GraphQL e coletamos 100 repositórios; na S02 "
        "escalamos para 1.000 e validamos os dados; na S03 analisamos e visualizamos as sete RQs.\n\n"
        "Objeto de estudo: os 1.000 repositórios públicos com mais estrelas no GitHub. "
        "Para linguagens mais populares (RQ05 e RQ07) adotamos o ranking do GitHub Octoverse 2025, "
        "mantido de forma consistente em todo o laboratório."
    ),
    "desafios": (
        "Rate limit da API GraphQL ao coletar mil repositórios — mitigado com leitura de rateLimit, "
        "retry e pausa quando o remaining fica baixo. Paginação da Search API (teto de 1.000 hits por "
        "query e nós nulos) exigiu janelas sucessivas por faixa de estrelas. O GitHub Projects v2 não "
        "expõe histórico de mudança de coluna via API, o que obrigou snapshots manuais ao fim de cada sprint. "
        "Ambiguidade em RQ04 entre pushedAt e updatedAt foi resolvida adotando pushedAt como métrica de "
        "atualização de código."
    ),
    "decisoes": (
        "WIP em Doing = 6 cartões (~2 por integrante no trio), para limitar paralelismo e forçar Review. "
        "RQ04 usa pushedAt, não updatedAt. RQ06 conta issues via GraphQL (sem PRs); repositórios com "
        "total_issues = 0 ficam fora da análise de closed_ratio. Outliers identificados pelo critério "
        "de Tukey (1,5 × IQR). Coleta feita com script próprio em Python, sem bibliotecas que encapsulam "
        "a API do GitHub."
    ),
    "etapas": (
        "S01 — Caio: infraestrutura GraphQL, coleta de 100 repos, snapshot do board. "
        "Henrique: validação das fatias de RQs. Jonas: hipóteses informais.\n\n"
        "S02 — Caio: paginação e coleta dos 1.000 repos, primeira versão do relatório, snapshot. "
        "Henrique: validação de consistência dos dados. Jonas: consolidação das hipóteses.\n\n"
        "S03 — Caio: análise RQ01–RQ03. Jonas: análise RQ04–RQ06. Henrique: cruzamento RQ07.\n\n"
        "Configuração do processo: colunas Backlog → To Do → Doing → Review → Done; WIP = 6; "
        "snapshots em snapshots/ (S01: 2026-08-13, S02: 2026-08-21, S03: 2026-08-27). "
        "Board: github.com/users/CaioVKodato/projects/6"
    ),
    "ferramentas": (
        "Python 3.12; requests e python-dotenv (cliente GraphQL próprio em src/github/); "
        "matplotlib e scipy (análise e gráficos); pandas (RQ04–RQ06). "
        "Processo: GitHub Projects v2 — github.com/users/CaioVKodato/projects/6. "
        "Repositório: github.com/CaioVKodato/lab-experimentacao-software"
    ),
    "inovacoes": (
        "RQ extra (S01): stars correlacionam com closed_ratio? "
        "Correlação de Spearman entre stars e closed_ratio nos 100 repos iniciais "
        "(89 válidos após excluir total_issues = 0). Resultado: Spearman ρ ≈ −0,05 — correlação nula. "
        "A hipótese se confirma: popularidade não prediz saúde do tracker. "
        "Documentação: docs/s01/caio/rq_extra_stars_vs_closed_ratio.md. "
        "Também adotamos paginação por janela de estrelas e snapshots GraphQL do Kanban como série "
        "temporal para os Labs 04 e 05."
    ),
    "coleta": (
        "Amostra final: 1.000 repositórios válidos em data/repositories_top1000.csv. "
        "Campos das RQs 01–05 e 07 completos (N = 1.000). "
        "RQ06: 957 repositórios com pelo menos uma issue; 43 excluídos (total_issues = 0). "
        "Outliers: 124 em merged_prs, 93 em releases, 196 em days_since_push (Tukey). "
        "Três snapshots do Kanban disponíveis (ago/2026)."
    ),
    "discussao": (
        "RQ01 — confirmada no centro da distribuição (mediana ~7,7 anos), com ressalva de 8,2% com menos de 1 ano.\n"
        "RQ02 — confirmada para o típico (mediana 768 PRs); 2% sem PR no GitHub.\n"
        "RQ03 — parcial: 27,7% sem nenhuma release.\n"
        "RQ04 — confirmada: mediana de 1 dia desde o último push.\n"
        "RQ05 — confirmada: domínio de Python, TypeScript e JavaScript.\n"
        "RQ06 — confirmada: mediana de 87,5% de issues fechadas.\n"
        "RQ07 — confirmada nas três métricas cruzadas por linguagem.\n\n"
        "Amostra mistura frameworks ativos com listas e tutoriais que viralizam sem ciclo de software "
        "tradicional — isso pesa nas RQs 02, 03 e 07. Ordenar por estrelas favorece ecossistemas grandes. "
        "A RQ extra reforça que stars não indicam manutenção de issues.\n\n"
        "A inovação não contradiz o enunciado: mostra que popularidade e closed_ratio são dimensões "
        "independentes, enquanto o restante descreve maturidade, atividade e engajamento."
    ),
    "conclusao": (
        "Repositórios muito estrelados costumam ser antigos, ativos e com comunidade grande, mas a amostra "
        "não é homogênea: listas e guias distorcem releases e PRs. As hipóteses do enunciado se confirmam "
        "no agregado, com RQ03 apenas parcial. Limitações: viés da Search API, mistura de tipos de projeto "
        "e métricas que dependem do uso do GitHub (PRs, Releases). Com mais tempo, segmentaríamos por tipo "
        "de repositório. A RQ extra e a série de snapshots valem expandir nos laboratórios seguintes."
    ),
}

METRICS_TABLE = [
    ["RQ01", "Idade do repositório", "Data da coleta − createdAt", "Dias", "GraphQL (createdAt)"],
    ["RQ02", "PRs aceitas", "pullRequests(states: MERGED).totalCount", "Contagem", "GraphQL"],
    ["RQ03", "Releases", "releases.totalCount", "Contagem", "GraphQL"],
    ["RQ04", "Última atualização", "Data da coleta − pushedAt", "Dias", "GraphQL (pushedAt)"],
    ["RQ05", "Linguagem primária", "primaryLanguage.name", "Categoria", "GraphQL"],
    ["RQ06", "Issues fechadas / total", "closed / (closed + open); exclui total = 0", "Proporção", "GraphQL issues"],
    ["RQ07", "Cruzamento por linguagem", "Medianas RQ02–04 por linguagem vs Octoverse", "—", "Script Python"],
    ["Extra", "Stars × closed_ratio", "Spearman entre stars e closed_ratio", "ρ", "Script Python (S01)"],
]

FIGURES = [
    ("RQ01 — Sistemas populares são maduros/antigos? Mediana: 2.826 dias (~7,7 anos).", DOCS / "s03" / "caio" / "analise_s03_rq01_idade.png", 14.0),
    ("RQ02 — Sistemas populares recebem muita contribuição externa? Mediana: 768 PRs aceitas.", None, 0),
    ("RQ03 — Sistemas populares lançam releases com frequência? Mediana: 40; 27,7% sem release.", None, 0),
    ("RQ04 — Sistemas populares são atualizados com frequência? Mediana: 1 dia desde o push.", DOCS / "s03" / "jonas" / "rq04_days_since_push_hist.png", 14.0),
    ("RQ05 — Sistemas populares usam linguagens populares? Python 22,7%, TypeScript 17,3%, JS 11,1%.", DOCS / "s03" / "jonas" / "rq05_language_bar.png", 14.0),
    ("RQ06 — Alto percentual de issues fechadas? Mediana: 87,5% (N = 957).", DOCS / "s03" / "jonas" / "rq06_closed_ratio_hist.png", 14.0),
    ("RQ07 — Linguagens populares têm mais PRs, releases e atualizações?", DOCS / "s03" / "henrique" / "rq07_por_linguagem.png", 16.5),
]


def fill_cover_table(doc: Document) -> None:
    t = doc.tables[0]
    data = {
        4: "Lab01 — Características de repositórios populares + Setup do Kanban",
        5: "Caio Victor Kodato Teixeira · Henrique Volponi · Jonas Martins",
        6: "github.com/CaioVKodato/lab-experimentacao-software · github.com/users/CaioVKodato/projects/6",
        7: "27/08/2026",
    }
    for idx, val in data.items():
        t.rows[idx].cells[1].text = val


def fill_metrics_table(doc: Document) -> None:
    t = doc.tables[1]
    while len(t.rows) < len(METRICS_TABLE) + 1:
        t.add_row()
    for i, row in enumerate(METRICS_TABLE, start=1):
        for j, val in enumerate(row):
            t.rows[i].cells[j].text = val


def remove_instructional_table(doc: Document) -> None:
    if len(doc.tables) > 2:
        doc.tables[2]._element.getparent().remove(doc.tables[2]._element)


def insert_figures_after(paragraph: Paragraph) -> None:
    anchor = paragraph
    for caption, path, width in FIGURES:
        anchor = insert_paragraph_after(anchor, caption)
        if path and path.is_file():
            anchor = insert_picture_after(anchor, path, width)


def process_document() -> None:
    shutil.copy2(TEMPLATE, OUT_DOCX)
    doc = Document(OUT_DOCX)

    fill_cover_table(doc)
    fill_metrics_table(doc)

    # Remover parágrafos indesejados
    for para in reversed(list(doc.paragraphs)):
        if should_skip(para.text.strip()):
            delete_paragraph(para)

    block_idx = 0
    figures_done = False

    for para in list(doc.paragraphs):
        text = para.text.strip()

        if text == "[Insira aqui os gráficos do grupo, um por RQ, cada um precedido da pergunta que ele responde]":
            replace_text(para, "")
            insert_figures_after(para)
            figures_done = True
            continue

        if text == "[Tabela ou linha do tempo com Sprint | Entregas | Responsável(is) | Issues (nº)]":
            replace_text(para, CONTENT["etapas"])
            continue

        if text == "[conteúdo do grupo — substituir este texto]":
            if block_idx < len(CONTENT_BLOCKS):
                replace_text(para, CONTENT[CONTENT_BLOCKS[block_idx]])
                block_idx += 1

    # Referências
    for para in doc.paragraphs:
        if para.text.strip().startswith("ZUSE, Horst"):
            replace_text(
                para,
                "ZUSE, Horst. A framework of software measurement. Walter de Gruyter, 2013.\n"
                "GITHUB. Octoverse 2025. Disponível em: "
                "https://github.blog/news-insights/octoverse/octoverse-a-new-developer-joins-github-every-second-as-ai-leads-typescript-to-1/\n"
                "GITHUB. GraphQL API documentation. Disponível em: https://docs.github.com/en/graphql",
            )
            break

    remove_instructional_table(doc)

    # Remover subtítulo de template na capa
    for para in list(doc.paragraphs):
        if "Modelo/Template de Relatório" in para.text:
            delete_paragraph(para)
            break

    fix_document_formatting(doc)

    doc.save(OUT_DOCX)
    print(f"DOCX gerado: {OUT_DOCX}")


if __name__ == "__main__":
    process_document()
